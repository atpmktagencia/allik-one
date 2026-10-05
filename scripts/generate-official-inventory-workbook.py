#!/usr/bin/env python3
"""Generate the official Allik One physical inventory workbook."""

from __future__ import annotations

import json
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "src/server/real-data/allik-oct-2026.json"
OUTPUT = ROOT / "docs/Allik_One_Planilha_Oficial_Estoque.xlsx"
MAX_ROWS = 500

# Cross-reference of the stock names with the vendor catalogs already registered.
# Ambiguous entries deliberately remain pending: a label/photo is required before renaming.
NOMENCLATURE = {
    "JDO-001-IM-ADEK": ("ADEK 600.000 UI/1 mL IM", "STIN-P007", "Vit. D3 600.000 UI + Vit. K2-MK7 1.300 mcg + Vit. A 50.000 UI + Vit. E 500 UI", "IM", "CONFIRMADO"),
    "JDO-002-IM-ANTI-INFLAMATORIO-E-ANTIOXID": ("Anti-inflamatório e Antioxidante — Protocolo EV 1.5", "ESS-EV1-5-P", "MSM 750 mg/5 mL + NAC 300 mg/2 mL + L-Carnitina 600 mg/2 mL + Complexo B sem B1 + SAME 200 mg/2 mL + Aminoácidos", "EV", "DIVERGÊNCIA DE VIA: SKU DIZ IM"),
    "JDO-003-IM-BOOSTER-ATP": ("Booster Energético — ATP 20 mg + L-Carnitina 600 mg + PQQ 5 mg/2 mL EV/IM", "STIN-P014", "ATP 20 mg + L-Carnitina 600 mg + PQQ 5 mg em 2 mL", "EV/IM", "CONFIRMADO"),
    "JDO-004-IM-BOOSTER-MOROSIL": ("Booster Morosil — concentração a confirmar — IM", "", "Composição e apresentação não localizadas nos catálogos cadastrados", "IM", "CONFIRMAR RÓTULO"),
    "JDO-005-IM-COENZIMA-Q10": ("Coenzima Q10 100 mg/2 mL IM", "STIN-P023", "Coenzima Q10 50 mg/mL; total 100 mg em 2 mL", "IM", "CONFIRMADO"),
    "JDO-006-IM-CURCUMINA": ("Curcumina 200 mg/2 mL IM", "STIN-P032", "Curcumina 100 mg/mL; total 200 mg em 2 mL", "IM", "CONFIRMADO"),
    "JDO-007-IM-INSULIN-SUPPORT": ("Insulin Support 5 mL IM", "STIN-P047", "Cromo Picolinato 100 mcg/mL + Inositol 100 mg/mL + Ácido Alfa-lipoico 10 mg/mL", "IM", "CONFIRMAR VOLUME COM FORNECEDOR"),
    "JDO-008-IM-L-GLUTATION": ("L-Glutationa 600 mg/5 mL EV/IM", "STIN-P055", "L-Glutationa 600 mg em 5 mL", "EV/IM", "CONFIRMADO"),
    "JDO-009-IM-NADH": ("NADH 50 mg — pó estéril para reconstituição — IM", "STIN-P075", "NADH 50 mg por frasco; volume de reconstituição não informado", "IM", "CONFIRMAR RECONSTITUIÇÃO"),
    "JDO-010-IM-NANDROLONA-DECANOATO": ("Nandrolona Decanoato — concentração a confirmar — IM", "", "Há apresentações diferentes; o produto não foi localizado no catálogo Stin cadastrado", "IM", "CONFIRMAR RÓTULO"),
    "JDO-011-EVIM-POOL-COGNITIVO-NEURO-HEALTH": ("Pool Cognitivo — composição a confirmar — EV/IM", "", "Fornecedor informado como Central Pharma; confirmar composição e volume no rótulo", "EV/IM", "CONFIRMAR RÓTULO"),
    "JDO-012-IM-RESVERATROL": ("Resveratrol 10 mg/1 mL IM/SC", "STIN-P089", "Resveratrol 10 mg em 1 mL", "IM/SC", "CONFIRMADO"),
    "JDO-013-IM-RESVERATROL-MICELAR": ("Nanomicelas de Resveratrol 10 mg/1 mL EV/IM", "ESS-P111", "Nanomicelas de Resveratrol 1%; 10 mg em 1 mL", "EV/IM", "CONFIRMADO"),
    "JDO-014-IM-TESTOSTERONA-CIPIONATO": ("Testosterona Cipionato — concentração a confirmar — IM", "", "Concentração e apresentação não confirmadas pelo fornecedor informado", "IM", "CONFIRMAR RÓTULO"),
    "JDO-015-IM-TESTOSTERONA-DURATESTON": ("Testosterona Durateston — composição/concentração a confirmar — IM", "", "Apresentação não localizada no catálogo Stin cadastrado", "IM", "CONFIRMAR RÓTULO"),
    "JDO-016-IM-TRIO-METILADOR": ("Trio Metilador 2 mL EV/IM", "STIN-P096", "Vit. B6 30 mg + Metilfolato 3.500 mcg + Vit. B12 2.500 mcg", "EV/IM", "CONFIRMADO"),
    "JDO-017-IM-TRIO-DETOX": ("Trio Detox — composição/concentração a confirmar — IM", "", "Composição não localizada nos catálogos cadastrados", "IM", "CONFIRMAR RÓTULO"),
    "JDO-018-IM-VITAMINA-B12": ("Vitamina B12 — forma e concentração a confirmar — IM", "", "Catálogo possui cianocobalamina e metilcobalamina em várias concentrações", "IM", "CONFIRMAR RÓTULO"),
    "JDO-019-IM-VITAMINA-D3": ("Vitamina D3 600.000 UI/1 mL IM", "STIN-P114", "Vitamina D3 600.000 UI/mL; ampola de 1 mL", "IM", "CONFIRMADO"),
    "JDO-020-EV-ALFA-LIPOICO-300": ("Ácido Alfa-lipoico 300 mg/5 mL EV", "STIN-P004", "Ácido Alfa-lipoico 60 mg/mL; total 300 mg em 5 mL", "EV/IM", "CONFIRMADO; SKU RESTRINGE EV"),
    "JDO-021-EV-ALFA-LIPOICO-600": ("Ácido Alfa-lipoico 600 mg/30 mL EV", "STIN-P002", "Ácido Alfa-lipoico 600 mg em 30 mL", "EV", "CONFIRMADO"),
    "JDO-022-EV-ANTI-INFLAMATORIO-E-ANTIOXID": ("Anti-inflamatório e Antioxidante — Protocolo EV 1.5", "ESS-EV1-5-P", "MSM 750 mg/5 mL + NAC 300 mg/2 mL + L-Carnitina 600 mg/2 mL + Complexo B sem B1 + SAME 200 mg/2 mL + Aminoácidos", "EV", "CONFIRMADO"),
    "JDO-023-EV-COGNICAO-E-MEMORIA-PLUS": ("Cognição e Memória Plus — Protocolo EV 3.10", "ESS-EV3-10-P", "Piracetam + L-Theanina + Inositol/Taurina + L-Fenilalanina + ATP + Aminoácidos + Alfa-GPC + Colina/L-Carnitina/B5", "EV", "CONFIRMADO"),
    "JDO-024-EV-DISBIOSE": ("Disbiose — composição a confirmar — EV", "", "Fornecedor informado como Stin; correspondência encontrada apenas no catálogo Essentia", "EV", "CONFIRMAR FORNECEDOR/RÓTULO"),
    "JDO-025-EV-ESTEATOSE-HEPATICA": ("Esteatose Hepática — composição a confirmar — EV", "", "Fornecedor informado como Stin; correspondência encontrada apenas no catálogo Essentia", "EV", "CONFIRMAR FORNECEDOR/RÓTULO"),
    "JDO-026-EV-FADIGA-INDISPOSICAO": ("Fadiga/Indisposição — Protocolo EV 2.1", "ESS-EV2-1-P", "NAC + Magnésio + Metilcobalamina + Complexo B sem B1 + D-Ribose + Taurina + Aminoácidos + Inositol", "EV", "CONFIRMADO"),
    "JDO-027-EV-FADIGA-MUSCULAR-E-HIPERTROFI": ("Fadiga Muscular e Hipertrofia — composição a confirmar — EV", "", "Composição exata não localizada com esse nome no catálogo Stin cadastrado", "EV", "CONFIRMAR RÓTULO"),
    "JDO-028-EV-FERRO": ("Ferro — princípio ativo, concentração e volume a confirmar — EV", "", "Não assumir carboximaltose sem conferir rótulo e fornecedor", "EV", "CONFIRMAR RÓTULO"),
    "JDO-029-EV-IMUNIDADE": ("Imunidade — Protocolo EV 1.1", "ESS-EV1-1-P", "Alanil Glutamina + Complexo B sem B1 + NAC + L-Glutationa + Minerais", "EV", "CONFIRMADO"),
    "JDO-030-EV-L-BAIBA": ("L-BAIBA 100 mg/1 mL EV", "ESS-P074", "Ácido L-β-aminoisobutírico 10%; 100 mg em 1 mL", "EV/IM", "CONFIRMADO; SKU RESTRINGE EV"),
    "JDO-031-EV-L-GLUTATION": ("L-Glutationa 600 mg/5 mL EV", "STIN-P055", "L-Glutationa 600 mg em 5 mL", "EV/IM", "CONFIRMADO; SKU RESTRINGE EV"),
    "JDO-032-Outra-LIDOCAINA": ("Lidocaína — concentração, volume e via a confirmar", "", "Existem várias concentrações e apresentações; conferir rótulo", "", "CONFIRMAR RÓTULO"),
    "JDO-033-EV-RESISTENCIA-INSULINICA": ("Resistência Insulínica — Protocolo EV 8.3", "ESS-EV8-3-P", "Minerais + L-Arginina + NAC + Cromo + Complexo B sem B1 + Magnésio + Vanádio", "EV", "CONFIRMADO"),
}

NAVY = "172033"
BLUE = "2563EB"
PALE_BLUE = "DBEAFE"
PALE_GREEN = "DCFCE7"
PALE_YELLOW = "FEF3C7"
PALE_RED = "FEE2E2"
WHITE = "FFFFFF"
GRAY = "64748B"
THIN = Side(style="thin", color="CBD5E1")


def title(ws, text: str, subtitle: str, last_column: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
    ws.cell(1, 1, text)
    ws.cell(1, 1).font = Font(size=18, bold=True, color=WHITE)
    ws.cell(1, 1).fill = PatternFill("solid", fgColor=NAVY)
    ws.cell(1, 1).alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 30
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_column)
    ws.cell(2, 1, subtitle)
    ws.cell(2, 1).font = Font(size=10, italic=True, color=GRAY)
    ws.cell(2, 1).alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = 32


def header(ws, row: int, labels: list[str]) -> None:
    for col, label in enumerate(labels, 1):
        cell = ws.cell(row, col, label)
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = Border(bottom=THIN)
    ws.row_dimensions[row].height = 34


def add_table(ws, ref: str, name: str) -> None:
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False
    )
    ws.add_table(table)


def add_list_validation(ws, column: str, formula: str, first=4, last=MAX_ROWS) -> None:
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    dv.error = "Selecione um valor da lista."
    dv.errorTitle = "Valor inválido"
    dv.prompt = "Escolha uma opção padronizada."
    dv.promptTitle = "Allik One"
    dv.showErrorMessage = True
    dv.showInputMessage = True
    ws.add_data_validation(dv)
    dv.add(f"{column}{first}:{column}{last}")


def build() -> None:
    data = json.loads(SOURCE.read_text(encoding="utf-8"))
    items = data["items"]
    prices = {(p["name"].strip().upper(), p["route"].strip().upper()): p["price"] for p in data["salePrices"]}

    wb = Workbook()
    wb.remove(wb.active)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True

    guide = wb.create_sheet("LEIA-ME")
    title(guide, "ALLIK ONE — PLANILHA OFICIAL DE ESTOQUE", "Versão 1.0 · Modelo para conferência física, saneamento cadastral e importação auditável", 6)
    instructions = [
        ("OBJETIVO", "Registrar o estoque real por unidade, local, produto e lote antes da atualização do Allik One."),
        ("1. PRODUTOS", "Revise SKU, nome, categoria, unidade de medida e estoque mínimo. Não reutilize um SKU para produtos diferentes."),
        ("2. CONTAGEM_FISICA", "Use uma linha para cada combinação de unidade + local + SKU + lote. Se houver dois lotes do mesmo produto, use duas linhas."),
        ("3. CAMPOS OBRIGATÓRIOS", "Unidade, local, SKU, lote, validade, status, quantidade física, custo unitário, fornecedor, data e responsável."),
        ("4. LOTE DESCONHECIDO", "Não invente o lote. Separe o item e informe SEM LOTE na observação; a liberação excepcional será decidida antes da importação."),
        ("5. VALIDADE", "Use uma data real. Não estimar validade. Produto vencido deve permanecer registrado e receber status BLOQUEADO."),
        ("6. CUSTO", "Informe o custo de aquisição por unidade de medida do cadastro, não o preço de venda. Use vírgula ou ponto conforme o Excel."),
        ("7. QUANTIDADE", "Conte na mesma unidade indicada em PRODUTOS. Não misture caixa, frasco, ampola e kit na mesma linha."),
        ("8. CONFERÊNCIA", "Depois do primeiro preenchimento, uma segunda pessoa deve revisar e preencher Segunda conferência por."),
        ("9. IMPORTAÇÃO", "Somente linhas com IMPORTAR = SIM e VALIDAÇÃO = OK serão consideradas. O sistema criará movimentos auditáveis; não haverá alteração direta de saldo."),
        ("10. SEGURANÇA", "Não inclua dados de pacientes, senhas, tokens ou credenciais nesta planilha."),
    ]
    guide.column_dimensions["A"].width = 24
    guide.column_dimensions["B"].width = 105
    for row, (label, text_value) in enumerate(instructions, 4):
        guide.cell(row, 1, label).font = Font(bold=True, color=NAVY)
        guide.cell(row, 1).fill = PatternFill("solid", fgColor=PALE_BLUE)
        guide.cell(row, 2, text_value).alignment = Alignment(wrap_text=True, vertical="top")
        guide.cell(row, 1).border = guide.cell(row, 2).border = Border(bottom=THIN)
        guide.row_dimensions[row].height = 38
    guide.sheet_view.showGridLines = False
    guide.freeze_panes = "A4"

    nomenclature = wb.create_sheet("NOMENCLATURA_ATIVOS")
    nomenclature_headers = ["SKU cadastrado", "Nome anterior", "Nome operacional padronizado", "Código no catálogo fornecedor", "Composição/concentração", "Via", "Fornecedor informado", "Situação", "Ação necessária"]
    title(nomenclature, "NOMENCLATURA OFICIAL DOS ATIVOS", "O nome padronizado descreve composição/concentração/apresentação e via. Não representa prescrição individual nem frequência de administração.", len(nomenclature_headers))
    header(nomenclature, 3, nomenclature_headers)
    for row, item in enumerate(items, 4):
        proposed, vendor_code, composition, route, status = NOMENCLATURE[item["sku"]]
        action = "Nenhuma" if status == "CONFIRMADO" else "Conferir foto do rótulo/embalagem antes de renomear ou importar"
        values = [item["sku"], item["name"], proposed, vendor_code, composition, route, item["supplier"], status, action]
        for col, value in enumerate(values, 1):
            nomenclature.cell(row, col, value)
        nomenclature.cell(row, 8).fill = PatternFill("solid", fgColor=PALE_GREEN if status == "CONFIRMADO" else PALE_YELLOW)
    nomenclature.freeze_panes = "A4"
    nomenclature.auto_filter.ref = f"A3:I{3 + len(items)}"
    for col, width in zip("ABCDEFGHI", [36,34,62,28,90,16,26,34,54]): nomenclature.column_dimensions[col].width = width
    for row in range(4, 4 + len(items)):
        for col in range(1, 10): nomenclature.cell(row, col).alignment = Alignment(wrap_text=True, vertical="top")
        nomenclature.row_dimensions[row].height = 58
    add_table(nomenclature, f"A3:I{3 + len(items)}", "NomenclaturaAtivos")

    products = wb.create_sheet("PRODUTOS")
    product_headers = ["SKU*", "Produto*", "Categoria*", "Unidade de medida*", "Estoque mínimo", "Preço de venda referência", "Rota", "Ativo?*", "Controla estoque?*", "Observações"]
    title(products, "CADASTRO OFICIAL DE PRODUTOS", "Revise o cadastro mestre. Campos com * são obrigatórios. Preço de venda é apenas referência comercial e não substitui o custo do lote.", len(product_headers))
    header(products, 3, product_headers)
    for row, item in enumerate(items, 4):
        key = (item["name"].strip().upper(), item["route"].strip().upper())
        proposed, _, _, _, status = NOMENCLATURE[item["sku"]]
        safe_name = proposed if status == "CONFIRMADO" else item["name"]
        values = [item["sku"], safe_name, "Injetáveis", item["unit"], item["minimum"], prices.get(key), item["route"], "SIM", "SIM", status if status != "CONFIRMADO" else item.get("notes")]
        for col, value in enumerate(values, 1):
            products.cell(row, col, value)
    for row in range(4 + len(items), MAX_ROWS + 1):
        products.cell(row, 8, "SIM")
        products.cell(row, 9, "SIM")
    add_list_validation(products, "D", "=LISTAS!$A$2:$A$8")
    add_list_validation(products, "G", "=LISTAS!$B$2:$B$6")
    add_list_validation(products, "H", "=LISTAS!$D$2:$D$3")
    add_list_validation(products, "I", "=LISTAS!$D$2:$D$3")
    products.freeze_panes = "A4"
    products.auto_filter.ref = f"A3:J{MAX_ROWS}"
    products.column_dimensions["A"].width = 35
    products.column_dimensions["B"].width = 46
    for col in "CDEFGHI": products.column_dimensions[col].width = 20
    products.column_dimensions["J"].width = 42
    for row in range(4, MAX_ROWS + 1):
        products.cell(row, 5).number_format = "0.000"
        products.cell(row, 6).number_format = 'R$ #,##0.00'
    add_table(products, f"A3:J{3 + len(items)}", "ProdutosOficiais")

    count = wb.create_sheet("CONTAGEM_FISICA")
    count_headers = [
        "IMPORTAR?*", "Unidade*", "Local de estoque*", "SKU*", "Produto (automático)", "Lote*", "Validade*",
        "Status do lote*", "Quantidade anterior", "Quantidade física*", "Diferença", "Unidade de medida (automático)",
        "Custo unitário*", "Fornecedor*", "Documento/NF", "Data da contagem*", "Contado por*", "Segunda conferência por",
        "Motivo/observações", "VALÍDAÇÃO"
    ]
    title(count, "CONTAGEM FÍSICA OFICIAL", "Uma linha por lote e local. Preencha as células amarelas; colunas azuis/cinzas são auxiliares. Não apague o histórico anterior.", len(count_headers))
    header(count, 3, count_headers)
    for idx, item in enumerate(items, 4):
        count.cell(idx, 1, "NÃO")
        count.cell(idx, 2, "Juazeiro do Norte")
        count.cell(idx, 3, "Estoque Central")
        count.cell(idx, 4, item["sku"])
        count.cell(idx, 5, f'=IFERROR(VLOOKUP(D{idx},PRODUTOS!$A:$J,2,FALSE),"")')
        count.cell(idx, 6, item.get("lot"))
        if item.get("expiresOn"):
            count.cell(idx, 7, datetime.strptime(item["expiresOn"], "%Y-%m-%d").date())
        count.cell(idx, 8, "DISPONÍVEL")
        count.cell(idx, 9, item["quantity"])
        count.cell(idx, 11, f'=IF(J{idx}="","",J{idx}-I{idx})')
        count.cell(idx, 12, f'=IFERROR(VLOOKUP(D{idx},PRODUTOS!$A:$J,4,FALSE),"")')
        count.cell(idx, 14, item.get("supplier"))
        count.cell(idx, 20, f'=IF(A{idx}<>"SIM","NÃO MARCADA",IF(COUNTA(B{idx}:H{idx},J{idx},M{idx}:N{idx},P{idx}:Q{idx})<12,"PREENCHER OBRIGATÓRIOS",IF(OR(J{idx}<0,M{idx}<0),"VALOR INVÁLIDO","OK")))')
    for row in range(4 + len(items), MAX_ROWS + 1):
        count.cell(row, 1, "NÃO")
        count.cell(row, 5, f'=IFERROR(VLOOKUP(D{row},PRODUTOS!$A:$J,2,FALSE),"")')
        count.cell(row, 8, "DISPONÍVEL")
        count.cell(row, 11, f'=IF(J{row}="","",J{row}-I{row})')
        count.cell(row, 12, f'=IFERROR(VLOOKUP(D{row},PRODUTOS!$A:$J,4,FALSE),"")')
        count.cell(row, 20, f'=IF(A{row}<>"SIM","NÃO MARCADA",IF(COUNTA(B{row}:H{row},J{row},M{row}:N{row},P{row}:Q{row})<12,"PREENCHER OBRIGATÓRIOS",IF(OR(J{row}<0,M{row}<0),"VALOR INVÁLIDO","OK")))')
    add_list_validation(count, "A", "=LISTAS!$D$2:$D$3")
    add_list_validation(count, "B", "=LISTAS!$E$2:$E$3")
    add_list_validation(count, "C", "=LISTAS!$F$2:$F$10")
    add_list_validation(count, "D", f"=PRODUTOS!$A$4:$A${MAX_ROWS}")
    add_list_validation(count, "H", "=LISTAS!$C$2:$C$4")
    count.freeze_panes = "A4"
    count.auto_filter.ref = f"A3:T{MAX_ROWS}"
    widths = [14,22,25,35,44,22,15,20,18,18,16,22,18,30,20,18,24,28,45,24]
    for i, width in enumerate(widths, 1): count.column_dimensions[chr(64+i) if i <= 26 else "A"].width = width
    yellow_columns = {1,2,3,4,6,7,8,10,13,14,15,16,17,18,19}
    for row in range(4, MAX_ROWS + 1):
        for col in range(1, 21):
            count.cell(row, col).border = Border(bottom=Side(style="hair", color="E2E8F0"))
            if col in yellow_columns:
                count.cell(row, col).fill = PatternFill("solid", fgColor=PALE_YELLOW)
        count.cell(row, 7).number_format = "dd/mm/yyyy"
        count.cell(row, 9).number_format = count.cell(row, 10).number_format = count.cell(row, 11).number_format = "0.000"
        count.cell(row, 13).number_format = 'R$ #,##0.0000'
        count.cell(row, 16).number_format = "dd/mm/yyyy"
    count.conditional_formatting.add(f"T4:T{MAX_ROWS}", FormulaRule(formula=['T4="OK"'], fill=PatternFill("solid", fgColor=PALE_GREEN)))
    count.conditional_formatting.add(f"T4:T{MAX_ROWS}", FormulaRule(formula=['AND(T4<>"OK",T4<>"NÃO MARCADA")'], fill=PatternFill("solid", fgColor=PALE_RED)))
    add_table(count, f"A3:T{3 + len(items)}", "ContagemOficial")

    suppliers = wb.create_sheet("FORNECEDORES")
    supplier_headers = ["Fornecedor*", "CNPJ", "Telefone", "E-mail", "Contato", "Ativo?*", "Observações"]
    title(suppliers, "CADASTRO DE FORNECEDORES", "Use o nome oficial e mantenha os dados de contato atualizados. Não inclua credenciais.", len(supplier_headers))
    header(suppliers, 3, supplier_headers)
    supplier_names = sorted({name.strip() for item in items for name in item.get("supplier", "").replace("/", ";").split(";") if name.strip()})
    known = {"Stin Pharma": "+55 11 2078-1800", "Essentia Pharma": "+55 48 8802-9876"}
    for row, name in enumerate(supplier_names, 4):
        suppliers.cell(row, 1, name)
        suppliers.cell(row, 3, known.get(name))
        suppliers.cell(row, 6, "SIM")
    add_list_validation(suppliers, "F", "=LISTAS!$D$2:$D$3")
    suppliers.freeze_panes = "A4"
    suppliers.auto_filter.ref = f"A3:G{MAX_ROWS}"
    for col, width in zip("ABCDEFG", [34,20,22,34,25,14,50]): suppliers.column_dimensions[col].width = width
    add_table(suppliers, f"A3:G{3 + len(supplier_names)}", "FornecedoresOficiais")

    lists = wb.create_sheet("LISTAS")
    list_columns = {
        "Unidades de medida": ["Caixa", "Kit", "Frasco", "Ampola", "Unidade", "Seringa", "Apresentação"],
        "Rotas": ["IM", "EV", "EV/IM", "SC", "Outra"],
        "Status do lote": ["DISPONÍVEL", "QUARENTENA", "BLOQUEADO"],
        "Sim/Não": ["SIM", "NÃO"],
        "Unidades": ["Fortaleza", "Juazeiro do Norte"],
        "Locais": ["Estoque Central", "Sala de Procedimentos", "Almoxarifado"],
    }
    for col, (label, values) in enumerate(list_columns.items(), 1):
        lists.cell(1, col, label).font = Font(bold=True, color=WHITE)
        lists.cell(1, col).fill = PatternFill("solid", fgColor=BLUE)
        for row, value in enumerate(values, 2): lists.cell(row, col, value)
        lists.column_dimensions[chr(64 + col)].width = 25
    lists.sheet_state = "hidden"

    for ws in wb.worksheets:
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_view.showGridLines = ws.title == "LISTAS"
        ws.auto_filter.showButton = True if ws.auto_filter.ref else False

    wb.save(OUTPUT)
    # Re-open to ensure the produced package is structurally valid.
    verified = load_workbook(OUTPUT, data_only=False)
    assert {"LEIA-ME", "NOMENCLATURA_ATIVOS", "PRODUTOS", "CONTAGEM_FISICA", "FORNECEDORES", "LISTAS"} <= set(verified.sheetnames)
    assert verified["CONTAGEM_FISICA"]["D4"].value == items[0]["sku"]
    print(OUTPUT)


if __name__ == "__main__":
    build()
